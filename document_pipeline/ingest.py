"""Document normalization / ingestion layer (Step 2).

Converts uploaded TXT/PDF/DOCX into a canonical Markdown representation plus
a metadata sidecar, deterministically (no LLM rewriting). Headings are
preserved where the source format actually carries structure (Word heading
styles), or inferred via regex heuristics for messy/unstructured text
(numbered Persian/English sections, "Chapter N"/"فصل ۱"/"بخش ۲", a positional
document title, and generic short-line subsection candidates). Bullet/numbered
lists are detected both from explicit markers and from runs of unmarked short
lines. Page boundaries are kept for PDFs. Falls back to plain paragraph blocks
when no structure is found, and never hard-splits a paragraph's concepts.
"""
import json
import os
import re
import shutil
from datetime import datetime, timezone
from typing import Dict, List, Optional, Tuple

from pypdf import PdfReader
from docx import Document as DocxDocument

from document_pipeline import llm_normalize

NORMALIZATION_VERSION = "v2"
CONVERTED_DIR = "./converted"
os.makedirs(CONVERTED_DIR, exist_ok=True)

# Step 2.5: optional LLM-assisted relabeling, off by default.
ENABLE_LLM_NORMALIZATION = os.getenv("ENABLE_LLM_NORMALIZATION", "false").lower() == "true"
LLM_NORMALIZATION_MODEL = os.getenv("LLM_NORMALIZATION_MODEL", "gemma3:12b")
LLM_NORMALIZATION_NUM_CTX = int(os.getenv("LLM_NORMALIZATION_NUM_CTX", "4096"))
LLM_NORMALIZATION_MAX_BATCHES = int(os.getenv("LLM_NORMALIZATION_MAX_BATCHES", "20"))

PERSIAN_DIGITS_TRANS = str.maketrans("۰۱۲۳۴۵۶۷۸۹", "0123456789")

# Explicit, low-false-positive structural signals.
CHAPTER_WORD_RE = re.compile(
    r"^(chapter|part|فصل|بخش)\s+([0-9۰-۹]+|[a-zA-Z]+)\b[:\.\-–—]?\s*(.*)$", re.IGNORECASE
)
NUMBERED_MARKER_RE = re.compile(r"^([0-9۰-۹]+)[\.\)]\s+(\S.*)$")
BULLET_MARKER_RE = re.compile(r"^([-*•])\s+(.*)$")

# Standalone boilerplate lines from ebook/PDF exports: a line that is entirely a
# bracketed note -- "[]", "[Cover for ...]", "[Image: ...]". These are cover/image
# alt-text placeholders, not real content, and must never become a title or heading.
BRACKET_ONLY_RE = re.compile(r"^\s*\[[^\]]*\]\s*$")

TRAILING_PUNCT_RE = re.compile(r"[.!۔:,،؛]\s*$")  # "?"/"؟" deliberately excluded: question-style headings are common
MAX_HEADING_LEN = 90
MAX_TITLE_LEN = 80
MAX_LIST_ITEM_LEN = 55  # deliberately tighter than MAX_HEADING_LEN: short fragments
# only -- a 60-90 char line is much more often a short transition *sentence*
# (often with multiple clauses) than a genuine list item.
MIN_IMPLICIT_LIST_RUN = 3

LOW_TEXT_DENSITY_CHARS_PER_PAGE = 50
OCR_REQUIRED_CHARS_PER_PAGE = 5

_DOCX_LIST_PREFIXES = ("list bullet", "list number", "list paragraph")


# ---------------------------------------------------------------------------
# Shared line/text classification helpers
# ---------------------------------------------------------------------------

def _has_latin_letters(s: str) -> bool:
    return bool(re.search(r"[A-Za-z]", s))


def _is_boilerplate_line(line: str) -> bool:
    """Standalone junk lines that should be dropped before structure detection:
    bracket-only placeholders ("[]", "[Cover ...]", "[Image: ...]"). Inline
    references like a trailing "[1]" inside a sentence are NOT matched (this only
    fires when the whole line is a single bracketed note)."""
    return bool(BRACKET_ONLY_RE.match(line or ""))


def _generic_heading_candidate(line: str) -> bool:
    """Weak, fuzzier heading signal: a short, unpunctuated, capitalized-looking
    standalone line. Only meaningful for Latin-script text (Persian has no
    letter case), so this naturally never fires on pure-Persian lines --
    that's intentional, not a bug: unmarked Persian subsection headings are a
    known limitation (see risks in the Step 2 report), and we'd rather miss
    them than over-detect ordinary short Persian sentences as headings."""
    line = line.strip()
    if not line or len(line) > MAX_HEADING_LEN or TRAILING_PUNCT_RE.search(line):
        return False
    if not _has_latin_letters(line):
        return False
    latin_words = [w for w in line.split() if re.search(r"[A-Za-z]", w)]
    if not latin_words:
        return False
    capitalized = sum(1 for w in latin_words if w[:1].isupper())
    return line.isupper() or capitalized >= max(1, len(latin_words) - 1)


def _is_title_candidate(line: str) -> bool:
    """Positional-only check (first block of a document): short, standalone,
    no terminal punctuation. Deliberately language-agnostic (no
    capitalization requirement) so Persian-only titles are detected too.
    Rejects boilerplate placeholders ("[]") and lines with no word characters
    so junk like "[]" is never promoted to the document title."""
    line = line.strip()
    if not line or len(line) > MAX_TITLE_LEN or TRAILING_PUNCT_RE.search(line):
        return False
    if _is_boilerplate_line(line) or not re.search(r"\w", line):
        return False
    return True


def _classify_marker_line(line: str) -> Tuple[Optional[str], Optional[re.Match]]:
    """Returns ('bullet'|'numbered', match) if line starts with an explicit
    list marker, else (None, None)."""
    m = BULLET_MARKER_RE.match(line)
    if m:
        return "bullet", m
    m = NUMBERED_MARKER_RE.match(line)
    if m:
        return "numbered", m
    return None, None


def _normalize_list_line(line: str) -> Optional[str]:
    """Normalizes a single explicit list-marker line to clean Markdown,
    preserving the original number for numbered items (not collapsed to 1.)."""
    kind, m = _classify_marker_line(line.strip())
    if kind == "bullet":
        return f"- {m.group(2).strip()}"
    if kind == "numbered":
        num = m.group(1).translate(PERSIAN_DIGITS_TRANS)
        return f"{num}. {m.group(2).strip()}"
    return None


def _classify_heading_line(line: str) -> Tuple[Optional[int], Optional[str]]:
    """Classifies a standalone line as a heading. Returns (level, strength)
    where strength is 'strong' (chapter-word / numbered-section -- low
    false-positive patterns) or 'weak' (generic short-line heuristic), or
    (None, None) if it's not a heading at all."""
    line = line.strip()
    if not line:
        return None, None
    # A real "Chapter 3"/"فصل ۲"/"Part II" heading is short, unpunctuated, and its
    # remainder (after the number) is empty or a Title -- not lowercase prose. Without
    # these guards, CHAPTER_WORD_RE also matched ordinary paragraphs that merely start
    # with one of those words ("Part of the neglect reflected ...", "Chapter 3 covered
    # how ...", including the ~80-char first line of a wrapped PDF paragraph) and
    # promoted body text to a section heading.
    m_chap = CHAPTER_WORD_RE.match(line)
    if m_chap and len(line) <= MAX_HEADING_LEN and not TRAILING_PUNCT_RE.search(line):
        rest = (m_chap.group(3) or "").strip()
        prose_continuation = bool(rest) and _has_latin_letters(rest) and rest[:1].islower()
        if not prose_continuation:
            return 2, "strong"
    nm = NUMBERED_MARKER_RE.match(line)
    if nm and len(line) <= MAX_HEADING_LEN and not TRAILING_PUNCT_RE.search(line):
        return 2, "strong"
    if _generic_heading_candidate(line):
        return 3, "weak"
    return None, None


def _is_short_list_candidate(line: str) -> bool:
    return bool(line) and len(line) <= MAX_LIST_ITEM_LEN


def _heading_stats(blocks: List[Tuple]) -> Dict[str, int]:
    strong = sum(1 for b in blocks if b[0] == "heading" and b[3] == "strong")
    weak = sum(1 for b in blocks if b[0] == "heading" and b[3] == "weak")
    return {"strong": strong, "weak": weak, "total": strong + weak}


def _structure_confidence(strong: int, weak: int, native: bool = False) -> str:
    if native or strong > 0:
        return "high"
    if weak > 0:
        return "medium"
    return "low"


def _promote_title(blocks: List[Tuple]) -> List[Tuple]:
    """If the very first block looks like a standalone document title,
    promote/upgrade it to a level-1 heading. Handles both a plain
    untouched paragraph AND a first line that the generic heuristic already
    (mis)classified as a lower-level heading -- e.g. a title containing a
    capitalized Latin product name still belongs at level 1, not 3. Only
    ever touches blocks[0]; never applied to a list block."""
    if not blocks:
        return blocks
    kind = blocks[0][0]
    if kind == "para":
        text = blocks[0][1]
        if _is_title_candidate(text):
            blocks[0] = ("heading", 1, text, "strong")
    elif kind == "heading":
        _, level, text, _strength = blocks[0]
        if level != 1 and _is_title_candidate(text):
            blocks[0] = ("heading", 1, text, "strong")
    return blocks


def _llm_meta_defaults() -> Dict:
    return {
        "llm_normalization_used": False,
        "llm_model": None,
        "llm_validation_status": "not_attempted",
        "llm_validation_warnings": [],
        "normalization_mode": "deterministic",
        "llm_batches_total": 0,
        "llm_batches_fallback_count": 0,
    }


def _maybe_llm_relabel(blocks: List[Tuple], structure_confidence: str, force: bool = False) -> Tuple[List[Tuple], Dict]:
    """Step 2.5: if enabled (ENABLE_LLM_NORMALIZATION=true) and either forced
    or the deterministic pass was unsure (structure_confidence not 'high'),
    asks llm_normalize to relabel ambiguous blocks. The original block text
    is NEVER replaced -- only the structural label can change. Returns
    (possibly-relabeled blocks, llm metadata fields to merge into the
    document's metadata dict)."""
    meta = _llm_meta_defaults()
    if not ENABLE_LLM_NORMALIZATION:
        return blocks, meta
    if not force and structure_confidence == "high":
        return blocks, meta

    new_blocks, info = llm_normalize.classify_document(
        blocks,
        model=LLM_NORMALIZATION_MODEL,
        num_ctx=LLM_NORMALIZATION_NUM_CTX,
        max_batches=LLM_NORMALIZATION_MAX_BATCHES,
        force=force,
    )
    meta["llm_batches_total"] = info.get("batches_total", 0)
    meta["llm_validation_warnings"] = info.get("warnings", [])
    if not info.get("used"):
        # Nothing to classify, or skipped for exceeding the batch-count cap.
        # normalization_mode stays "deterministic"; status records why.
        meta["llm_validation_status"] = info.get("status", "not_attempted")
        return blocks, meta

    meta["llm_normalization_used"] = True
    meta["llm_model"] = LLM_NORMALIZATION_MODEL
    meta["llm_validation_status"] = info["status"]
    meta["llm_batches_fallback_count"] = info["batches_fallback"]
    meta["normalization_mode"] = "fallback" if info["status"] == "failed" else "llm_assisted"
    return new_blocks, meta


# ---------------------------------------------------------------------------
# Shared tokenizer core, used by TXT, PDF-per-page, and DOCX alike.
#
# `entries` is a list of (text, forced) pairs where `forced` is either None
# (needs heuristic classification, exactly like a plain text line) or an
# already-known classification: ("heading", level, "strong") -- e.g. from a
# real Word heading style -- or ("list", normalized_text) -- e.g. from a real
# Word list style. This lets DOCX's authoritative styling and TXT/PDF's
# regex heuristics share the exact same marker/heading/implicit-list
# disambiguation logic instead of duplicating (and drifting from) it.
# ---------------------------------------------------------------------------

def _tokenize(entries: List[Tuple[str, Optional[Tuple]]], detect_implicit_lists: bool) -> List[Tuple]:
    n = len(entries)
    blocks: List[Tuple] = []
    para_buf: List[str] = []

    def flush_para():
        if not para_buf:
            return
        blocks.append(("para", " ".join(para_buf).strip()))
        para_buf.clear()

    def is_plain(idx: int) -> bool:
        return idx < n and entries[idx][1] is None

    def is_special_plain_line(t: str) -> bool:
        """True if a plain line would itself be intercepted as a marker/
        heading by the main loop -- used to stop an implicit-list run."""
        if _classify_marker_line(t)[0] is not None:
            return True
        return _classify_heading_line(t)[0] is not None

    i = 0
    while i < n:
        text, forced = entries[i]
        if not text:
            flush_para()
            i += 1
            continue
        if forced is not None:
            flush_para()
            if forced[0] == "heading":
                blocks.append(("heading", forced[1], text, forced[2]))
            else:
                blocks.append(("list", forced[1]))
            i += 1
            continue

        kind, _m = _classify_marker_line(text)
        if kind == "bullet":
            flush_para()
            items = []
            while is_plain(i):
                k2, m2 = _classify_marker_line(entries[i][0])
                if k2 != "bullet":
                    break
                items.append(f"- {m2.group(2).strip()}")
                i += 1
            blocks.append(("list", "\n".join(items)))
            continue
        if kind == "numbered":
            next_kind = _classify_marker_line(entries[i + 1][0])[0] if is_plain(i + 1) else None
            if next_kind == "numbered":
                flush_para()
                items = []
                while is_plain(i):
                    k2, m2 = _classify_marker_line(entries[i][0])
                    if k2 != "numbered":
                        break
                    num = m2.group(1).translate(PERSIAN_DIGITS_TRANS)
                    items.append(f"{num}. {m2.group(2).strip()}")
                    i += 1
                blocks.append(("list", "\n".join(items)))
                continue
            # Isolated numbered line (not part of a run): falls through to
            # the heading classifier below, which applies its own
            # length/punctuation safety checks before accepting it.

        level, strength = _classify_heading_line(text)
        if level:
            flush_para()
            blocks.append(("heading", level, text, strength))
            i += 1
            continue

        if detect_implicit_lists and _is_short_list_candidate(text):
            # Look ahead for a run of 3+ consecutive short, otherwise-plain
            # lines -- an unmarked list embedded between ordinary sentences
            # (e.g. soft-broken lines within one paragraph), not just a
            # uniformly-short whole paragraph. Carve it out as its own list
            # block even if longer sentences precede/follow it.
            run = []
            j = i
            while is_plain(j):
                t2 = entries[j][0]
                if not t2 or is_special_plain_line(t2) or not _is_short_list_candidate(t2):
                    break
                run.append(t2)
                j += 1
            if len(run) >= MIN_IMPLICIT_LIST_RUN:
                flush_para()
                blocks.append(("list", "\n".join(f"- {l}" for l in run)))
                i = j
                continue

        para_buf.append(text)
        i += 1
    flush_para()
    return blocks


def _lines_to_blocks(text: str, detect_implicit_lists: bool = True) -> List[Tuple]:
    """Tokenize raw text (TXT, or one PDF page) into ('heading', level,
    title, strength) / ('list', text) / ('para', text) blocks.

    detect_implicit_lists controls whether runs of 3+ unmarked short lines
    get promoted to a bullet list. This is safe for TXT (newlines are
    usually real line breaks) but is deliberately disabled for PDF page text,
    where pypdf's line breaks are often just visual word-wrap artifacts that
    would otherwise misclassify ordinary wrapped paragraphs as lists."""
    raw_lines = [l.strip() for l in text.replace("\r\n", "\n").replace("\r", "\n").split("\n")]
    # Drop standalone boilerplate (cover/image placeholders) by turning them into
    # blank separators so they neither become content nor a heading/title.
    raw_lines = ["" if _is_boilerplate_line(l) else l for l in raw_lines]
    entries = [(l, None) for l in raw_lines]
    return _tokenize(entries, detect_implicit_lists)


def _blocks_to_markdown(blocks: List[Tuple]) -> str:
    md_parts = []
    for b in blocks:
        if b[0] == "heading":
            _, level, title, _strength = b
            md_parts.append(("#" * level) + " " + title)
        else:
            md_parts.append(b[1])
    return "\n\n".join(md_parts)


# ---------------------------------------------------------------------------
# TXT
# ---------------------------------------------------------------------------

def normalize_txt(text: str, force_llm_normalization: bool = False) -> Dict:
    blocks = _lines_to_blocks(text, detect_implicit_lists=True)
    blocks = _promote_title(blocks)
    stats = _heading_stats(blocks)
    confidence = _structure_confidence(stats["strong"], stats["weak"])

    blocks, llm_meta = _maybe_llm_relabel(blocks, confidence, force=force_llm_normalization)
    if llm_meta["llm_normalization_used"]:
        stats = _heading_stats(blocks)
        confidence = _structure_confidence(stats["strong"], stats["weak"])

    markdown_text = _blocks_to_markdown(blocks)
    warning = None
    if confidence == "low":
        warning = "weak_structure_no_headings_detected"
    meta = {
        "extraction_method": "raw-txt",
        "structure_confidence": confidence,
        "page_count": None,
        "heading_count": stats["total"],
        "char_count": len(markdown_text),
        "extraction_quality_warning": warning,
    }
    meta.update(llm_meta)
    return {"markdown_text": markdown_text, "meta": meta}


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------

def normalize_docx(file_stream, force_llm_normalization: bool = False) -> Dict:
    doc = DocxDocument(file_stream)
    entries: List[Tuple[str, Optional[Tuple]]] = []
    style_heading_count = 0

    for p in doc.paragraphs:
        # python-docx renders soft line breaks (Shift+Enter) within a single
        # Word paragraph as embedded "\n" in p.text -- split those into
        # separate candidate lines, the same unit size as a TXT line, so
        # e.g. an unmarked short-line list typed with soft breaks inside one
        # paragraph still gets per-line classification instead of being
        # treated as one opaque multi-line blob.
        raw = p.text.replace("\r\n", "\n").replace("\r", "\n")
        sub_lines = [
            l.strip() for l in raw.split("\n")
            if l.strip() and not _is_boilerplate_line(l.strip())
        ]
        if not sub_lines:
            continue
        style = (p.style.name if p.style is not None else "") or ""
        style_lower = style.lower()

        if style_lower.startswith("heading") or style_lower == "title":
            if style_lower == "title":
                level = 1
            else:
                digits = re.search(r"(\d+)", style)
                level = min(int(digits.group(1)), 3) if digits else 1
            entries.append((" ".join(sub_lines), ("heading", level, "strong")))
            style_heading_count += 1
        elif style_lower.startswith(_DOCX_LIST_PREFIXES):
            full = " ".join(sub_lines)
            entries.append((full, ("list", _normalize_list_line(full) or f"- {full}")))
        else:
            # No forced classification: the shared tokenizer will apply the
            # same heuristics used for plain text, so messy student docs
            # (everything left in "Normal" style) still get structure.
            for sl in sub_lines:
                entries.append((sl, None))

    blocks = _tokenize(entries, detect_implicit_lists=True)
    blocks = _promote_title(blocks)
    stats = _heading_stats(blocks)
    confidence = _structure_confidence(stats["strong"], stats["weak"], native=style_heading_count > 0)

    blocks, llm_meta = _maybe_llm_relabel(blocks, confidence, force=force_llm_normalization)
    if llm_meta["llm_normalization_used"]:
        stats = _heading_stats(blocks)
        confidence = _structure_confidence(stats["strong"], stats["weak"], native=style_heading_count > 0)

    markdown_text = _blocks_to_markdown(blocks)
    warning = None
    if confidence == "low":
        warning = "weak_structure_no_headings_detected"
    meta = {
        "extraction_method": "python-docx",
        "structure_confidence": confidence,
        "page_count": None,  # python-docx exposes no reliable page-layout info
        "heading_count": stats["total"],
        "char_count": len(markdown_text),
        "extraction_quality_warning": warning,
    }
    meta.update(llm_meta)
    return {"markdown_text": markdown_text, "meta": meta}


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def _ocr_backend_available() -> Tuple[bool, str]:
    """Read-only check (no installs) for a usable OCR backend. Used only when
    ENABLE_OCR_FALLBACK is explicitly turned on; otherwise never invoked."""
    if shutil.which("tesseract") is None:
        return False, "tesseract executable not found on PATH"
    try:
        import pytesseract  # noqa: F401
    except ImportError:
        return False, "pytesseract not installed"
    try:
        import pdf2image  # noqa: F401
    except ImportError:
        return False, "pdf2image not installed"
    return True, ""


def normalize_pdf(file_stream, enable_ocr_fallback: Optional[bool] = None, force_llm_normalization: bool = False) -> Dict:
    if enable_ocr_fallback is None:
        enable_ocr_fallback = os.getenv("ENABLE_OCR_FALLBACK", "false").lower() == "true"

    reader = PdfReader(file_stream)
    page_count = len(reader.pages)
    total_chars = 0
    page_blocks_list: List[List[Tuple]] = []
    all_blocks: List[Tuple] = []

    for i, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        total_chars += len(text.strip())
        page_blocks = _lines_to_blocks(text, detect_implicit_lists=False)
        if i == 0:
            page_blocks = _promote_title(page_blocks)
        page_blocks_list.append(page_blocks)
        all_blocks.extend(page_blocks)

    stats = _heading_stats(all_blocks)
    confidence = _structure_confidence(stats["strong"], stats["weak"])

    all_blocks, llm_meta = _maybe_llm_relabel(all_blocks, confidence, force=force_llm_normalization)
    if llm_meta["llm_normalization_used"]:
        # Split the relabeled flat sequence back into per-page chunks (using
        # each page's original block count) so page markers stay correctly
        # interleaved during rendering below.
        lengths = [len(pb) for pb in page_blocks_list]
        page_blocks_list = []
        cursor = 0
        for length in lengths:
            page_blocks_list.append(all_blocks[cursor:cursor + length])
            cursor += length
        stats = _heading_stats(all_blocks)
        confidence = _structure_confidence(stats["strong"], stats["weak"])

    md_blocks = []
    for i, page_blocks in enumerate(page_blocks_list):
        md_blocks.append(f"<!-- page:{i + 1} -->")
        page_md = _blocks_to_markdown(page_blocks)
        if page_md:
            md_blocks.append(page_md)
    markdown_text = "\n\n".join(md_blocks)
    avg_chars_per_page = (total_chars / page_count) if page_count else 0

    warning = None
    ocr_meta = {}
    if avg_chars_per_page < OCR_REQUIRED_CHARS_PER_PAGE:
        warning = "ocr_required"
    elif avg_chars_per_page < LOW_TEXT_DENSITY_CHARS_PER_PAGE:
        warning = "low_text_density_possible_scanned_pdf"
    elif confidence == "low":
        warning = "weak_structure_no_headings_detected"

    if warning in ("ocr_required", "low_text_density_possible_scanned_pdf"):
        ocr_meta["ocr_fallback_enabled"] = enable_ocr_fallback
        if enable_ocr_fallback:
            available, reason = _ocr_backend_available()
            ocr_meta["ocr_backend_available"] = available
            if not available:
                ocr_meta["ocr_backend_missing_reason"] = reason
            # Deterministic, no-LLM scaffolding only: actually rendering pages
            # to images and running tesseract is not implemented here because
            # no OCR backend is installed locally (see ingest.OCR_SETUP_NOTES).
            # This flag/path exists so it can be wired up later without
            # touching the upload route again.

    meta = {
        "extraction_method": "pypdf",
        "structure_confidence": confidence,
        "page_count": page_count,
        "heading_count": stats["total"],
        "char_count": len(markdown_text),
        "extraction_quality_warning": warning,
    }
    meta.update(ocr_meta)
    meta.update(llm_meta)
    return {"markdown_text": markdown_text, "meta": meta}


OCR_SETUP_NOTES = (
    "No local OCR backend found (tesseract executable, pytesseract, pdf2image, "
    "ocrmypdf all absent). Lightest Windows-friendly setup if/when approved: "
    "1) install the Tesseract-OCR Windows installer (UB-Mannheim build, ~80-100MB, "
    "includes the 'fas' Persian language pack option) and add it to PATH; "
    "2) pip install pytesseract (~50KB) + pdf2image (~30KB); "
    "3) pdf2image additionally needs the Poppler Windows binaries on PATH "
    "(~10-20MB zip, no installer, just unzip+PATH). Total download ~100-150MB, "
    "all one-time. Nothing is installed by this module on its own."
)


# ---------------------------------------------------------------------------
# Dispatch / IO
# ---------------------------------------------------------------------------

def normalize_document(filename: str, file_stream, ext: str, force_llm_normalization: bool = False) -> Dict:
    """Dispatch by extension. Returns {"markdown_text": str, "meta": {...}}.
    force_llm_normalization bypasses the structure_confidence trigger (still
    requires ENABLE_LLM_NORMALIZATION=true to actually call the LLM) -- meant
    for manual testing on a specific document, not for production traffic."""
    ext = ext.lower()
    if ext == ".txt":
        text = file_stream.read()
        if isinstance(text, bytes):
            text = text.decode("utf-8")
        return normalize_txt(text, force_llm_normalization=force_llm_normalization)
    if ext == ".docx":
        return normalize_docx(file_stream, force_llm_normalization=force_llm_normalization)
    if ext == ".pdf":
        return normalize_pdf(file_stream, force_llm_normalization=force_llm_normalization)
    raise ValueError(f"Unsupported file type: {ext}")


def write_normalized(document_id: str, markdown_text: str, metadata: Dict) -> Tuple[str, str]:
    """Writes converted/{document_id}.md and .metadata.json. Returns their
    relative paths. Callers should set metadata["normalized_md_path"] (e.g.
    via paths_for()) BEFORE calling this, so the field is present in the
    written JSON, not just in the caller's in-memory dict."""
    md_path, meta_path = paths_for(document_id)
    with open(md_path, "w", encoding="utf-8") as f:
        f.write(markdown_text)
    with open(meta_path, "w", encoding="utf-8") as f:
        json.dump(metadata, f, ensure_ascii=False, indent=2)
    return md_path, meta_path


def paths_for(document_id: str) -> Tuple[str, str]:
    md_path = os.path.join(CONVERTED_DIR, f"{document_id}.md")
    meta_path = os.path.join(CONVERTED_DIR, f"{document_id}.metadata.json")
    return md_path, meta_path


def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()
